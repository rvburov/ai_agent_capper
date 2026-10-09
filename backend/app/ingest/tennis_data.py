"""Загрузка архива tennis-data.co.uk: даты матчей, кэфы букмекеров и турниры, которых еще
нет в архиве Сакманна.

    python -m app.ingest.tennis_data [--years 2013-2026] [--dir ПАПКА] [--force]

Файл года — {год}/{год}.xlsx (основной тур ATP, с 2013 года). Лист один, первая строка —
заголовки; колонки читаются по именам, неизвестные пропускаются, колонок кэфов может
не быть. Строки группируются в турнир по (год, ATP), с группой работаем так:

1. Привязка к турниру Сакманна. Кандидаты — турниры source = sackmann, начатые не раньше
   чем за 16 дней до даты строки и не позже чем через 3 дня после нее (турниры Большого
   шлема начинаются в воскресенье, а у Сакманна дата — понедельник). Строка «голосует»
   за кандидата, если в нем есть матч той же пары игроков (имя «Фамилия И.» против
   имен и алиасов игроков матча, app.ingest.names). Турнир группы — кандидат, за
   которого проголосовала хотя бы половина строк и больше, чем за любого другого.
   Строки группы сопоставляются только с его матчами; пара, сыгранная дважды (круговой
   этап и финал), уточняется по раунду. Найденному матчу ставится день (match_date) и
   кэфы, турниру — зал, город, покрытие и категория, если у Сакманна их нет или
   категория общая (atp). Счет, статус и победителя не трогаем: главный источник —
   Сакманн. Строка, пары которой в выбранном турнире нет, идет только в отчет
   (match_not_found), своего матча для нее нет: чаще всего это расхождение имен, и свой
   матч стал бы дублем матча Сакманна (Эло и бэктест посчитали бы его дважды).
2. Привязки нет (архив Сакманна отстает или турнира там нет): свой турнир (source =
   tennis_data, source_ref = «{год}-{ATP}») и свои матчи (source_ref — хеш года, ATP,
   метки раунда и имен). В source_ref — номер круга из метки («2nd Round» → 2), а не
   код R32/R64: код зависит от оценки сетки по строкам группы, а она растет, пока
   турнир идет. Код раунда у матча обновляется при каждом разборе. Игроки ищутся среди
   всех игроков базы по ключам «фамилия + инициалы»; из нескольких подходящих берется
   единственный, у кого есть матчи в пределах двух лет от даты строки; не нашлось
   никого — создается игрок «И. Фамилия» (позже его «усыновит» загрузка Сакманна).
   Несколько подходящих — строка в отчет, матч не создается. Строка, пара которой
   нашлась в турнире Сакманна, но турнир группы не выбран, тоже идет в отчет, а не в
   свой матч: иначе матч попал бы в базу дважды.
3. Свои матчи турнира сверяются с файлом. Матч, которого в файле больше нет (строку
   исправили: перепутаны победитель и проигравший, другое имя), удаляется вместе с
   кэфами, если на него ничего не ссылается (ставки, прогнозы, досье, статистика, кэфы
   не из архива). Иначе, а также если ему может соответствовать строка, которая в этот
   раз не сопоставилась, матч остается, а в отчет пишется предупреждение.

Игроки старого бота, у которых нет sackmann_id и все матчи из его истории (legacy), в
поиске не участвуют: их сливает с настоящими игроками привязка старого бота, а матч
tennis-data с таким игроком помешал бы ей.

Тезок tennis-data различает двумя буквами имени: «Kuznetsov An.» и «Kuznetsov Al.».
Игрок, у которого имя начинается с других двух букв, такому имени не подходит ни в
поиске, ни в турнире Сакманна, даже по первому инициалу (TdName.contradicts).

Кэфы — вид archive, рынок ML, taken_at — полночь дня матча (UTC): selection 1 — кэф игрока
player1 матча в базе. Колонки: PS* — pinnacle, B365* — bet365, Max* — max, Avg* — avg,
BFE* — betfair_ex. Повторная загрузка обновляет кэф, а не добавляет строку. Когда сняты
кэфы (закрытие или раньше), сайт не сообщает.

Каждый год — своя транзакция. Неизменившийся файл (sha256 как в прошлом успешном
запуске) не разбирается без --force. Файла текущего года нет — не ошибка, а отметка в
freshness. Нет файла прошлого года на сайте (там они есть всегда) или не нашлось ни
одного файла из запрошенных (пустая или не та папка) — ошибка: остальные годы
загружаются, а запуск помечается неудачным. Ответ 403 или веб-страница вместо файла
(сайт закрыт Cloudflare от программ) — ошибка с подсказкой скачать файлы в браузере и
загрузить их из папки.
"""

import argparse
import asyncio
import hashlib
import sys
from collections import Counter, defaultdict
from collections.abc import Callable, Iterable, Sequence
from dataclasses import dataclass, field
from datetime import UTC, date, datetime, time, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path
from typing import Any, Literal, get_args

from openpyxl import load_workbook
from openpyxl.worksheet._read_only import ReadOnlyWorksheet
from sqlalchemy import (
    BigInteger,
    Boolean,
    ColumnElement,
    Row,
    and_,
    bindparam,
    delete,
    exists,
    func,
    literal_column,
    or_,
    select,
    text,
    update,
)
from sqlalchemy.dialects.postgresql import ARRAY, insert
from sqlalchemy.exc import DBAPIError
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import get_settings
from app.core.i18n import t
from app.db.base import Base
from app.db.models import SURFACES, Match, OddsSnapshot, Player, Tournament
from app.db.session import create_engine, create_sessionmaker
from app.ingest import rounds
from app.ingest.fetch import DataSource, FetchError
from app.ingest.names import fold, name_keys, parse_tennis_data_name
from app.ingest.runs import (
    IngestError,
    IngestReport,
    ingestion_run,
    last_ok_file_hashes,
    update_lock,
)

SOURCE = "tennis_data"
MARKET = "ML"
ODDS_KIND = "archive"
# Страница, где файлы можно скачать в браузере, если сайт не отдает их программе
MANUAL_PAGE = "http://www.tennis-data.co.uk/alldata.php"
# Окно кандидатов: турнир Сакманна начат не позже чем через 3 дня после даты строки и
# не раньше чем за 16 дней до нее (две недели турнира Большого шлема и запас)
DAYS_BEFORE_START = 3
DAYS_AFTER_START = 16
# Из нескольких игроков с подходящим именем берется тот, кто играл в пределах 2 лет
ACTIVE_YEARS = 2
# Строк в одном INSERT: 2000 × 7 параметров меньше предела asyncpg (32 767)
BATCH = 2000
# Строку заголовков ищем среди первых строк листа
HEADER_ROWS = 10
ODDS_MIN = Decimal("1.01")
ODDS_MAX = Decimal("999.999")
ODDS_STEP = Decimal("0.001")
# День 0 в датах Excel (с поправкой на несуществующее 29 февраля 1900 года)
EXCEL_EPOCH = date(1899, 12, 30)
EXCEL_MAX = 2_958_466
# Ключ pg_advisory_xact_lock: два загрузчика tennis-data не пишут один год одновременно
LOCK_KEY = 0x7464617461_0001
SCORE_LIMIT = 64

# Без этих колонок файл не разобрать
REQUIRED = ("atp", "tournament", "date", "round", "winner", "loser", "comment")
# Префикс колонок кэфов (PSW/PSL …) → букмекер в odds_snapshots.bookmaker
BOOKMAKERS = (
    ("ps", "pinnacle"),
    ("b365", "bet365"),
    ("max", "max"),
    ("avg", "avg"),
    ("bfe", "betfair_ex"),
)
# Series (без пробелов, нижний регистр) → tournaments.category. International и
# International Gold — названия 250 и 500 до 2009 года
SERIES = {
    "grandslam": "grand_slam",
    "masters1000": "masters",
    "masters": "masters",
    "atp500": "500",
    "internationalgold": "500",
    "atp250": "250",
    "international": "250",
    "masterscup": "finals",
    "atpfinals": "finals",
}
# Категория турнира Сакманна, которую tennis-data уточняет (250 или 500)
GENERAL_CATEGORY = "atp"
# Comment → статус матча; прочее (Sched и т. п.) — матч не сыгран, строка пропускается
STATUSES = {
    "completed": "finished",
    "retired": "retired",
    "disqualified": "retired",
    "walkover": "walkover",
}

UnmatchedKind = Literal[
    "match_not_found",
    "match_ambiguous",
    "tournament_ambiguous",
    "in_sackmann",
    "player_ambiguous",
    "same_player",
    "duplicate_row",
]
UNMATCHED_KINDS: tuple[str, ...] = get_args(UnmatchedKind)
# Строка запроса к matches (поля по имени)
type _DbRow = Row[*tuple[Any, ...]]

# --- разбор файла ---


def clean_text(value: object) -> str | None:
    """Текст ячейки без лишних пробелов; пустая ячейка — None. 10.0 → «10»."""
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    cleaned = " ".join(str(value).split())
    return cleaned or None


def to_date(value: object) -> date | None:
    """Дата из ячейки: datetime, date, число Excel (45306 → 2024-01-15) или строка
    «15/01/2024», «2024-01-15». Иначе None."""
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, int | float) and not isinstance(value, bool):
        if not 1 <= value < EXCEL_MAX:
            return None
        return EXCEL_EPOCH + timedelta(days=int(value))
    if isinstance(value, str):
        raw = value.strip()
        for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d.%m.%Y", "%Y-%m-%d %H:%M:%S"):
            try:
                return datetime.strptime(raw, fmt).date()
            except ValueError:
                continue
    return None


def to_int(value: object) -> int | None:
    """Целое из ячейки: 6, 6.0, «6». Иначе None."""
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    number: float
    if isinstance(value, float):
        number = value
    elif isinstance(value, str):
        try:
            number = float(value.strip().replace(",", "."))
        except ValueError:
            return None
    else:
        return None
    return int(number) if number.is_integer() else None


def to_odds(value: object) -> Decimal | None:
    """Кэф из ячейки с тремя знаками. Пустая ячейка — None; не число, меньше 1.01 или
    больше 999.999 — ValueError."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, bool) or not isinstance(value, int | float | str | Decimal):
        raise ValueError(value)
    try:
        number = Decimal(str(value).strip().replace(",", "."))
    except InvalidOperation:
        raise ValueError(value) from None
    if not number.is_finite() or number < ODDS_MIN or number > ODDS_MAX:
        raise ValueError(value)
    return number.quantize(ODDS_STEP, ROUND_HALF_UP)


def series_category(series: str | None) -> str | None:
    """Series tennis-data → категория турнира: Grand Slam → grand_slam, Masters 1000 →
    masters, ATP500 → 500, ATP250 → 250, Masters Cup → finals. Неизвестная — None."""
    if not series:
        return None
    return SERIES.get("".join(series.split()).casefold())


def comment_key(comment: str | None) -> str:
    """Comment без лишних пробелов и регистра: « Retired » → «retired»."""
    return " ".join((comment or "").split()).casefold()


def match_status(comment: str | None) -> str | None:
    """Comment → статус матча: Completed → finished, Retired и Disqualified → retired,
    Walkover → walkover. Прочее — None (матч не сыгран)."""
    return STATUSES.get(comment_key(comment))


def build_score(sets: Sequence[tuple[int, int]], comment: str) -> str | None:
    """Счет в формате Сакманна по сетам «6-4 3-6 7-6». Тай-брейков у tennis-data нет,
    поэтому «7-6» без очков. Отказ — « RET», дисквалификация — « DEF», неявка — «W/O»."""
    key = comment_key(comment)
    if key == "walkover":
        return "W/O"
    parts = [f"{won}-{lost}" for won, lost in sets]
    if key == "retired":
        parts.append("RET")
    elif key == "disqualified":
        parts.append("DEF")
    return " ".join(parts)[:SCORE_LIMIT] or None


def player_display_name(raw: str) -> str:
    """Имя нового игрока из «Фамилия И.»: «Del Potro J.M.» → «J.M. Del Potro»."""
    tokens = raw.split()
    surname, initials = parse_tennis_data_name(raw)
    if not initials:
        return " ".join(tokens)
    cut = next(
        (n for n in range(len(tokens), 0, -1) if fold(" ".join(tokens[:n])) == surname),
        len(tokens),
    )
    return " ".join(tokens[cut:] + tokens[:cut])


def round_key(label: str) -> str:
    """Раунд для source_ref своего матча без размера сетки: «2nd Round» → «2»,
    Quarterfinals и Quarter-finals → «QF». Код R32/R64 сюда не годится: он зависит от
    оценки сетки, а она меняется, пока турнир идет."""
    number = rounds.round_number(label)
    if number is not None:
        return str(number)
    try:
        return rounds.tennis_data_round(label, rounds.MAX_DRAW)
    except ValueError:
        return fold(label)


def own_match_ref(year: int, atp: int, round_label: str, winner: str, loser: str) -> str:
    """source_ref своего матча tennis-data: одинаков между запусками, пока не поменялись
    год, номер турнира, метка раунда и имена (размер сетки не входит)."""
    raw = f"{year}|{atp}|{round_key(round_label)}|{fold(winner)}|{fold(loser)}"
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()[:32]


@dataclass
class TdRow:
    """Строка файла tennis-data после разбора."""

    line: int
    atp: int
    day: date
    tournament: str
    location: str | None
    series: str | None
    indoor: bool | None
    surface: str | None
    round_label: str
    best_of: int | None
    winner: str
    loser: str
    status: str
    score: str | None
    # (букмекер, сторона: 1 — Winner, 2 — Loser) → кэф
    odds: dict[tuple[str, int], Decimal] = field(default_factory=dict)

    @property
    def pinnacle(self) -> bool:
        return ("pinnacle", 1) in self.odds and ("pinnacle", 2) in self.odds

    def example(self, path: str, **extra: Any) -> dict[str, Any]:
        """Поля строки для примера в отчете."""
        return {
            "file": path,
            "line": self.line,
            "date": self.day,
            "tournament": self.tournament,
            "round": self.round_label,
            "winner": self.winner,
            "loser": self.loser,
            **extra,
        }


@dataclass
class Sheet:
    """Лист файла: колонки (имя в нижнем регистре → номер) и непустые строки данных."""

    columns: dict[str, int]
    rows: list[tuple[int, tuple[object, ...]]]


def _header_key(value: object) -> str | None:
    cleaned = clean_text(value)
    return cleaned.casefold() if cleaned else None


def _looks_like_html(content: bytes) -> bool:
    head = content[:2048].lstrip().lower()
    return head.startswith((b"<!doctype", b"<html", b"<head", b"<body")) or b"<html" in head


def read_sheet(content: bytes, path: str) -> Sheet:
    """Первый лист файла xlsx: строка заголовков и строки данных. Файл не xlsx, нет
    заголовков или нужных колонок — IngestError."""
    if not content.startswith(b"PK"):
        if _looks_like_html(content):
            raise IngestError("ingest_tennis_data.errors.not_xlsx", path=path, page=MANUAL_PAGE)
        raise IngestError("ingest_tennis_data.errors.bad_file", path=path, error="not xlsx")
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        # Битый архив, не та структура: BadZipFile, KeyError, InvalidFileException и др.
        error = type(exc).__name__
        raise IngestError("ingest_tennis_data.errors.bad_file", path=path, error=error) from None
    columns: dict[str, int] | None = None
    rows: list[tuple[int, tuple[object, ...]]] = []
    try:
        sheet = workbook.worksheets[0]
        if isinstance(sheet, ReadOnlyWorksheet):
            # Размер листа в файле бывает записан неверно: читаем до конца
            sheet.reset_dimensions()
        for number, values in enumerate(sheet.iter_rows(values_only=True), start=1):
            if columns is None:
                found = {
                    key: index
                    for index, key in reversed(list(enumerate(map(_header_key, values))))
                    if key
                }
                if "winner" in found and "loser" in found:
                    columns = found
                elif number >= HEADER_ROWS:
                    break
                continue
            if any(value is not None and value != "" for value in values):
                rows.append((number, tuple(values)))
    except IndexError:
        columns = None
    except Exception as exc:
        error = type(exc).__name__
        raise IngestError("ingest_tennis_data.errors.bad_file", path=path, error=error) from None
    finally:
        workbook.close()
    if columns is None:
        raise IngestError("ingest_tennis_data.errors.no_header", path=path)
    missing = [name for name in REQUIRED if name not in columns]
    if missing:
        raise IngestError(
            "ingest_tennis_data.errors.missing_columns", path=path, columns=", ".join(missing)
        )
    return Sheet(columns, rows)


def parse_rows(sheet: Sheet, *, path: str, year: int, report: IngestReport) -> list[TdRow]:
    """Строки листа → TdRow. Строка без имен, номера турнира, раунда или даты и матч,
    который не сыгран (Comment не Completed, Retired, Disqualified, Walkover), —
    пропуск с пометкой в отчете. Негодный кэф пропускается, строка остается."""
    columns = sheet.columns

    def cell(values: tuple[object, ...], name: str) -> object:
        index = columns.get(name)
        return values[index] if index is not None and index < len(values) else None

    odds_columns = [
        (bookmaker, side, f"{prefix}{suffix}")
        for prefix, bookmaker in BOOKMAKERS
        for side, suffix in ((1, "w"), (2, "l"))
        if f"{prefix}{suffix}" in columns
    ]
    parsed: list[TdRow] = []
    for line, values in sheet.rows:
        report.rows_read += 1
        atp = to_int(cell(values, "atp"))
        winner = clean_text(cell(values, "winner"))
        loser = clean_text(cell(values, "loser"))
        round_label = clean_text(cell(values, "round"))
        tournament = clean_text(cell(values, "tournament"))
        if (
            atp is None
            or not winner
            or not loser
            or not round_label
            or not tournament
            or fold(winner) == fold(loser)
        ):
            report.skip("bad_row")
            report.add_warning("ingest_tennis_data.warnings.bad_row", path=path, line=line)
            continue
        raw_date = cell(values, "date")
        day = to_date(raw_date)
        if day is None or abs(day.year - year) > 1:
            report.skip("bad_date")
            report.add_warning(
                "ingest_tennis_data.warnings.bad_date", path=path, line=line, value=raw_date
            )
            continue
        comment = clean_text(cell(values, "comment"))
        status = match_status(comment)
        if status is None:
            report.skip("status")
            report.add_warning(
                "ingest_tennis_data.warnings.status_skipped",
                path=path,
                line=line,
                comment=comment or "",
            )
            continue
        sets: list[tuple[int, int]] = []
        for number in range(1, 6):
            won, lost = to_int(cell(values, f"w{number}")), to_int(cell(values, f"l{number}"))
            if won is not None and lost is not None:
                sets.append((won, lost))
        surface = (clean_text(cell(values, "surface")) or "").casefold()
        court = (clean_text(cell(values, "court")) or "").casefold()
        best_of = to_int(cell(values, "best of"))
        row = TdRow(
            line=line,
            atp=atp,
            day=day,
            tournament=tournament,
            location=clean_text(cell(values, "location")),
            series=clean_text(cell(values, "series")),
            indoor={"indoor": True, "outdoor": False}.get(court),
            surface=surface if surface in SURFACES else None,
            round_label=round_label,
            best_of=best_of if best_of in (3, 5) else None,
            winner=winner,
            loser=loser,
            status=status,
            score=build_score(sets, comment or ""),
        )
        for bookmaker, side, column in odds_columns:
            raw = cell(values, column)
            try:
                odds = to_odds(raw)
            except ValueError:
                report.counts["odds.invalid"] += 1
                report.add_warning(
                    "ingest_tennis_data.warnings.odds_invalid",
                    path=path,
                    line=line,
                    column=column.upper(),
                    value=raw,
                )
                continue
            if odds is not None:
                row.odds[(bookmaker, side)] = odds
        parsed.append(row)
    return parsed


def parse_file(content: bytes, *, path: str, year: int, report: IngestReport) -> list[TdRow]:
    """Файл xlsx → строки (синхронно: вызывается в отдельном потоке)."""
    return parse_rows(read_sheet(content, path), path=path, year=year, report=report)


# --- имена и игроки ---


@dataclass(frozen=True)
class TdName:
    """Имя tennis-data («Фамилия И.») и его ключи для сравнения с игроками базы.

    strong — ключи с инициалами («sinner j», «del potro jm», «ramos a»), weak — только
    фамилия. tiers — ключи для поиска по всей базе, от точного к общему: полные
    инициалы, первый инициал, прочие ключи с инициалом (короткая форма составной
    фамилии: «Ramos-Vinolas A.» → «ramos a»). prefix — начало имени, которым tennis-data
    различает тезок («Kuznetsov An.» → «an»), иначе пусто."""

    text: str
    keys: frozenset[str]
    strong: frozenset[str]
    weak: frozenset[str]
    tiers: tuple[frozenset[str], ...]
    prefix: str

    @classmethod
    def parse(cls, text: str) -> "TdName":
        surname, initials = parse_tennis_data_name(text)
        keys = frozenset(name_keys(text))
        forms = {surname, surname.replace(" ", "")}
        weak = frozenset(forms)
        if initials:
            exact = frozenset({fold(text), *(f"{form} {initials}" for form in forms)})
            first = frozenset(f"{form} {initials[0]}" for form in forms)
            # Ключ вида «o c» из «O'Connell C.» слишком общий для поиска по всей базе
            rest = frozenset(key for key in keys - weak - exact - first if len(key.split()[0]) >= 3)
            tiers = tuple(tier for tier in (exact, first, rest) if tier)
        else:
            tiers = (weak,)
        return cls(text, keys, keys - weak, weak, tiers, _name_prefix(text, initials))

    def contradicts(self, player_keys: frozenset[str]) -> bool:
        """Игрок — другой тезка: в файле «Kuznetsov Al.», а у игрока имя начинается с
        других двух букв (ключ «kuznetsov an» у Andrey Kuznetsov). Без этой проверки
        «Al.» совпал бы с ним по первому инициалу."""
        if not self.prefix:
            return False
        two = self.prefix[:2]
        same = {f"{form} {two}" for form in self.weak}
        if not same.isdisjoint(player_keys):
            return False
        for key in player_keys:
            for form in self.weak:
                rest = key.removeprefix(f"{form} ")
                if rest != key and len(rest) == 2 and rest.isalpha() and rest[0] == two[0]:
                    return True
        return False

    def level(self, player_keys: frozenset[str]) -> int:
        """Насколько игрок похож: 2 — фамилия и инициал, 1 — только фамилия, 0 — нет
        (или это другой тезка, см. contradicts)."""
        if self.contradicts(player_keys):
            return 0
        if not self.strong.isdisjoint(player_keys):
            return 2
        if not self.weak.isdisjoint(player_keys):
            return 1
        return 0


def _name_prefix(text: str, initials: str) -> str:
    """Две-три первые буквы имени одним словом («An.», «Ale.») — так tennis-data различает
    тезок. «J.M.» и «JM» — инициалы двух имен (Juan Martin), а не начало одного: пусто.
    Поэтому «Kwon S.W.» не противоречит игроку Soonwoo Kwon (ключ «kwon so»)."""
    tokens = text.split()
    letters = tokens[-1].rstrip(".") if tokens else ""
    if (
        len(initials) >= 2
        and letters.isalpha()
        and not letters.isupper()
        and fold(letters) == initials
    ):
        return initials
    return ""


def _legacy_only() -> ColumnElement[bool]:
    """Игрок есть только в истории старого бота: нет sackmann_id, все матчи legacy."""
    plays = or_(Match.player1_id == Player.id, Match.player2_id == Player.id)
    return and_(
        Player.sackmann_id.is_(None),
        exists().where(plays, Match.source == "legacy"),
        ~exists().where(plays, Match.source != "legacy"),
    )


_ACTIVE_YEARS_SQL = text(
    "SELECT p.id, array_agg(DISTINCT EXTRACT(YEAR FROM COALESCE("
    "  m.match_date, t.start_date, (m.scheduled_at AT TIME ZONE 'UTC')::date))::int)"
    " FROM unnest(CAST(:ids AS bigint[])) AS p(id)"
    " JOIN matches m ON m.player1_id = p.id OR m.player2_id = p.id"
    " LEFT JOIN tournaments t ON t.id = m.tournament_id"
    " GROUP BY p.id"
).bindparams(bindparam("ids", type_=ARRAY(BigInteger)))


@dataclass
class _Ambiguous:
    """Под имя подходят несколько игроков."""

    ids: list[int]


class PlayerIndex:
    """Ключи имен всех игроков базы (имя и алиасы, app.ingest.names.name_keys): поиск
    по имени tennis-data без запроса к базе на каждую строку."""

    def __init__(self) -> None:
        self.keys: dict[int, frozenset[str]] = {}
        self.by_key: defaultdict[str, set[int]] = defaultdict(set)
        # Игроки, среди которых ищется имя своего матча tennis-data
        self.searchable: set[int] = set()
        # Годы, в которые у игрока есть матчи (загружаются по мере надобности)
        self.years: dict[int, set[int]] = {}
        self._names: dict[str, TdName] = {}

    def add(self, player_id: int, names: Iterable[str], *, searchable: bool) -> None:
        if player_id in self.keys:
            return
        keys = frozenset(key for name in names if name for key in name_keys(name))
        self.keys[player_id] = keys
        for key in keys:
            self.by_key[key].add(player_id)
        if searchable:
            self.searchable.add(player_id)

    async def load(self, db: AsyncSession, ids: Iterable[int] | None = None) -> None:
        """Игроки базы (все или только ids)."""
        stmt = select(Player.id, Player.name, Player.aliases, _legacy_only())
        if ids is not None:
            stmt = stmt.where(Player.id.in_(list(ids)))
        for player_id, name, aliases, legacy_only in await db.execute(stmt):
            self.add(player_id, [name, *(aliases or [])], searchable=not legacy_only)

    def name(self, text_: str) -> TdName:
        if text_ not in self._names:
            self._names[text_] = TdName.parse(text_)
        return self._names[text_]

    def player_keys(self, player_id: int) -> frozenset[str]:
        return self.keys.get(player_id, frozenset())

    def candidates(self, name: TdName) -> set[int]:
        """Игроки, подходящие под имя, по первому уровню ключей, где кто-то нашелся.
        Тезка с другим началом имени (TdName.contradicts) не подходит."""
        for tier in name.tiers:
            found = {
                pid
                for key in tier
                for pid in self.by_key.get(key, ())
                if pid in self.searchable and not name.contradicts(self.keys[pid])
            }
            if found:
                return found
        return set()

    async def active_years(self, db: AsyncSession, ids: Iterable[int]) -> dict[int, set[int]]:
        missing = [pid for pid in ids if pid not in self.years]
        if missing:
            for pid in missing:
                self.years[pid] = set()
            rows = await db.execute(_ACTIVE_YEARS_SQL, {"ids": missing})
            for pid, years in rows:
                self.years[pid] = {int(year) for year in years if year is not None}
        return {pid: self.years[pid] for pid in ids}

    def played(self, player_id: int, year: int) -> None:
        """Отметить матч игрока в году (для выбора из нескольких подходящих)."""
        if player_id in self.years:
            self.years[player_id].add(year)

    async def resolve(self, db: AsyncSession, name: TdName, day: date) -> int | _Ambiguous | None:
        """Игрок по имени: id; None — никто не подходит; _Ambiguous — несколько."""
        found = self.candidates(name)
        if len(found) == 1:
            return next(iter(found))
        if not found:
            return None
        years = await self.active_years(db, found)
        near = sorted(
            pid for pid in found if any(abs(year - day.year) <= ACTIVE_YEARS for year in years[pid])
        )
        if len(near) == 1:
            return near[0]
        return _Ambiguous(sorted(found))


# --- турниры и матчи Сакманна ---


@dataclass
class _SackMatch:
    id: int
    round: str
    player1_id: int
    player2_id: int
    winner_id: int | None
    match_date: date | None


@dataclass
class _Found:
    """Матч кандидата с той же парой: score — сумма похожести двух имен (2–4),
    flipped — победитель tennis-data в базе записан вторым игроком."""

    score: int
    match: _SackMatch
    flipped: bool


@dataclass
class _Candidate:
    """Турнир Сакманна с матчами и указателем «ключ имени → матчи»."""

    tournament: Tournament
    start: date
    matches: list[_SackMatch]
    by_key: dict[str, set[int]]
    draw: int

    @classmethod
    def build(
        cls, tournament: Tournament, matches: list[_SackMatch], index: PlayerIndex
    ) -> "_Candidate":
        assert tournament.start_date is not None
        by_key: defaultdict[str, set[int]] = defaultdict(set)
        for number, match in enumerate(matches):
            for player_id in (match.player1_id, match.player2_id):
                for key in index.player_keys(player_id):
                    by_key[key].add(number)
        sizes = [
            int(match.round[1:])
            for match in matches
            if match.round.startswith("R") and match.round[1:].isdigit()
        ]
        draw = max(sizes, default=rounds.MIN_DRAW)
        return cls(tournament, tournament.start_date, matches, dict(by_key), draw)

    def covers(self, day: date) -> bool:
        return (
            self.start - timedelta(days=DAYS_BEFORE_START)
            <= day
            <= self.start + timedelta(days=DAYS_AFTER_START)
        )

    def find(
        self, winner: TdName, loser: TdName, index: PlayerIndex, claimed: set[int]
    ) -> list[_Found]:
        """Матчи этого турнира с парой (winner, loser) в любом порядке, кроме уже
        занятых другими строками."""
        around_winner = {n for key in winner.keys for n in self.by_key.get(key, ())}
        around_loser = {n for key in loser.keys for n in self.by_key.get(key, ())}
        found: list[_Found] = []
        for number in sorted(around_winner & around_loser):
            match = self.matches[number]
            if match.id in claimed:
                continue
            first = index.player_keys(match.player1_id)
            second = index.player_keys(match.player2_id)
            options = [
                (winner.level(first), loser.level(second), False),
                (winner.level(second), loser.level(first), True),
            ]
            best = max(
                (option for option in options if min(option[0], option[1]) >= 1),
                key=lambda option: option[0] + option[1],
                default=None,
            )
            if best is not None:
                found.append(_Found(best[0] + best[1], match, best[2]))
        return found


def _common[T](values: Iterable[T | None]) -> T | None:
    """Самое частое непустое значение."""
    counter = Counter(value for value in values if value is not None)
    return counter.most_common(1)[0][0] if counter else None


@dataclass
class _GroupInfo:
    """Сведения о турнире по строкам группы."""

    name: str
    city: str | None
    series: str | None
    category: str | None
    indoor: bool | None
    surface: str | None
    best_of: int | None
    start: date

    @classmethod
    def of(cls, rows: Sequence[TdRow]) -> "_GroupInfo":
        series = _common(row.series for row in rows)
        best = [row.best_of for row in rows if row.best_of is not None]
        return cls(
            name=(_common(row.tournament for row in rows) or "")[:200],
            city=(_common(row.location for row in rows) or "")[:100] or None,
            series=series,
            category=series_category(series),
            indoor=_common(row.indoor for row in rows),
            surface=_common(row.surface for row in rows),
            best_of=max(best) if best else None,
            start=min(row.day for row in rows),
        )


@dataclass
class _Plan:
    """Свой матч tennis-data, который нужно записать."""

    row: TdRow
    tournament_id: int
    round: str
    winner_id: int
    loser_id: int
    ref: str

    @property
    def key(self) -> tuple[int, str, int, int]:
        return (self.tournament_id, self.round, self.winner_id, self.loser_id)

    def values(self) -> dict[str, Any]:
        return {
            "tournament_id": self.tournament_id,
            "round": self.round,
            "player1_id": self.winner_id,
            "player2_id": self.loser_id,
            "winner_id": self.winner_id,
            "match_date": self.row.day,
            "status": self.row.status,
            "score": self.row.score,
            "best_of": self.row.best_of,
        }


@dataclass
class _YearStats:
    rows: int = 0
    pinnacle: int = 0
    linked: int = 0
    own: int = 0
    unmatched: int = 0
    first_date: date | None = None
    last_date: date | None = None


def _batches[T](items: Sequence[T], size: int = BATCH) -> Iterable[Sequence[T]]:
    for start in range(0, len(items), size):
        yield items[start : start + size]


class _YearLoader:
    """Запись одного файла (года) в одной транзакции."""

    def __init__(
        self,
        db: AsyncSession,
        *,
        year: int,
        path: str,
        rows: list[TdRow],
        index: PlayerIndex,
        report: IngestReport,
    ) -> None:
        self.db = db
        self.year = year
        self.path = path
        self.rows = rows
        self.index = index
        self.report = report
        self.stats = _YearStats()
        # Матчи Сакманна, уже сопоставленные строкам этого файла
        self.claimed: set[int] = set()
        self.match_dates: dict[int, date] = {}
        # (матч, selection, букмекер) → (кэф, день матча)
        self.odds: dict[tuple[int, int, str], tuple[Decimal, date]] = {}
        self.plans: list[_Plan] = []
        self.own_tournaments: dict[str, Tournament] = {}
        self.own_match_count: Counter[int] = Counter()
        # Свои турниры групп этого файла → строки группы, не ставшие матчем: матчи
        # турнира сверяются с файлом, а матч, которому может соответствовать такая
        # строка, не удаляется
        self.own_groups: dict[int, list[TdRow]] = {}

    def _count(self, key: str, created: bool) -> None:
        self.report.counts[key] += 1
        if created:
            self.report.created += 1
        else:
            self.report.updated += 1

    def _unmatched(self, kind: UnmatchedKind, row: TdRow, **extra: Any) -> None:
        self.stats.unmatched += 1
        self.report.add_unmatched(kind, **row.example(self.path, **extra))

    async def run(self) -> _YearStats:
        self.stats.rows = len(self.rows)
        self.stats.pinnacle = sum(1 for row in self.rows if row.pinnacle)
        if not self.rows:
            return self.stats
        days = [row.day for row in self.rows]
        self.stats.first_date, self.stats.last_date = min(days), max(days)
        await self.db.execute(text("SELECT pg_advisory_xact_lock(:key)"), {"key": LOCK_KEY})
        candidates = await self._load_candidates(
            min(days) - timedelta(days=DAYS_AFTER_START),
            max(days) + timedelta(days=DAYS_BEFORE_START),
        )
        await self._load_own_tournaments()
        groups: defaultdict[int, list[TdRow]] = defaultdict(list)
        for row in self.rows:
            groups[row.atp].append(row)
        # По порядку турниров: матч, занятый турниром прошлой недели, не достанется
        # турниру следующей, где та же пара встретилась снова
        ordered = sorted(groups.items(), key=lambda item: (min(r.day for r in item[1]), item[0]))
        for atp, group in ordered:
            near = [cand for cand in candidates if any(cand.covers(row.day) for row in group)]
            chosen = self._choose(group, near)
            info = _GroupInfo.of(group)
            if isinstance(chosen, _Candidate):
                self._link(atp, group, info, chosen)
            else:
                await self._own(atp, group, info, near, ambiguous=chosen)
        await self._write_match_dates()
        await self._write_own_matches()
        await self._write_odds()
        return self.stats

    # --- загрузка ---

    async def _load_candidates(self, start_from: date, start_to: date) -> list[_Candidate]:
        tournaments = list(
            (
                await self.db.scalars(
                    select(Tournament)
                    .where(
                        Tournament.source == "sackmann",
                        Tournament.start_date.between(start_from, start_to),
                    )
                    .order_by(Tournament.start_date, Tournament.id)
                )
            ).all()
        )
        if not tournaments:
            return []
        by_tournament: defaultdict[int, list[_SackMatch]] = defaultdict(list)
        rows = await self.db.execute(
            select(
                Match.id,
                Match.tournament_id,
                Match.round,
                Match.player1_id,
                Match.player2_id,
                Match.winner_id,
                Match.match_date,
            )
            .where(Match.tournament_id.in_([tour.id for tour in tournaments]))
            .order_by(Match.id)
        )
        for mid, tid, round_code, p1, p2, winner, match_date in rows:
            if tid is not None:
                by_tournament[tid].append(_SackMatch(mid, round_code, p1, p2, winner, match_date))
        players = {
            pid
            for matches in by_tournament.values()
            for match in matches
            for pid in (match.player1_id, match.player2_id)
        }
        unknown = players - self.index.keys.keys()
        if unknown:
            await self.index.load(self.db, unknown)
        return [_Candidate.build(tour, by_tournament[tour.id], self.index) for tour in tournaments]

    async def _load_own_tournaments(self) -> None:
        found = await self.db.scalars(
            select(Tournament).where(
                Tournament.source == SOURCE, Tournament.source_ref.like(f"{self.year}-%")
            )
        )
        self.own_tournaments = {tour.source_ref or "": tour for tour in found}
        if self.own_tournaments:
            counts = await self.db.execute(
                select(Match.tournament_id, func.count())
                .where(Match.tournament_id.in_([t.id for t in self.own_tournaments.values()]))
                .group_by(Match.tournament_id)
            )
            self.own_match_count.update({tid: int(n) for tid, n in counts if tid is not None})

    # --- привязка к Сакманну ---

    def _choose(self, group: list[TdRow], near: list[_Candidate]) -> _Candidate | bool:
        """Турнир Сакманна для группы. False — нет подходящего, True — двое поровну."""
        if not near:
            return False
        votes: Counter[int] = Counter()
        for row in group:
            winner, loser = self.index.name(row.winner), self.index.name(row.loser)
            for cand in near:
                if cand.covers(row.day) and cand.find(winner, loser, self.index, self.claimed):
                    votes[cand.tournament.id] += 1
        if not votes:
            return False
        ranked = votes.most_common()
        best_id, best = ranked[0]
        if best * 2 < len(group):
            return False
        if len(ranked) > 1 and ranked[1][1] == best:
            return True
        return next(cand for cand in near if cand.tournament.id == best_id)

    def _link(self, atp: int, group: list[TdRow], info: _GroupInfo, cand: _Candidate) -> None:
        """Строки группы → матчи турнира Сакманна. Строка без пары в нем — только в отчет
        (match_not_found), свой матч не создается: см. docstring модуля, п. 1."""
        if self._enrich(cand.tournament, info):
            self._count("tournaments.updated", created=False)
        stale = self.own_tournaments.get(f"{self.year}-{atp}")
        if stale is not None and self.own_match_count[stale.id]:
            self.report.add_warning(
                "ingest_tennis_data.warnings.stale_tournament",
                path=self.path,
                tournament=info.name,
                ref=stale.source_ref,
            )
        for row in group:
            winner, loser = self.index.name(row.winner), self.index.name(row.loser)
            found = cand.find(winner, loser, self.index, self.claimed)
            if not found:
                self._unmatched("match_not_found", row, sackmann_tournament=cand.tournament.id)
                continue
            top = max(item.score for item in found)
            chosen = [item for item in found if item.score == top]
            if len(chosen) > 1:
                try:
                    code: str | None = rounds.tennis_data_round(row.round_label, cand.draw)
                except ValueError:
                    code = None
                chosen = [item for item in chosen if item.match.round == code]
            if len(chosen) != 1:
                ids = sorted(item.match.id for item in found if item.score == top)
                self._unmatched("match_ambiguous", row, match_ids=ids)
                continue
            self._apply_link(row, chosen[0])

    def _enrich(self, tournament: Tournament, info: _GroupInfo) -> bool:
        """Зал, город, покрытие и категория турнира Сакманна, если там пусто; категория
        atp (250 или 500 — неизвестно) уточняется."""
        changed = False
        if tournament.indoor is None and info.indoor is not None:
            tournament.indoor = info.indoor
            changed = True
        if not tournament.city and info.city:
            tournament.city = info.city
            changed = True
        if tournament.surface is None and info.surface is not None:
            tournament.surface = info.surface
            changed = True
        if (
            info.category is not None
            and tournament.category in (None, GENERAL_CATEGORY)
            and tournament.category != info.category
        ):
            tournament.category = info.category
            changed = True
        return changed

    def _apply_link(self, row: TdRow, found: _Found) -> None:
        match = found.match
        self.claimed.add(match.id)
        self.stats.linked += 1
        self.report.counts["rows.linked"] += 1
        if match.match_date != row.day:
            self.match_dates[match.id] = row.day
            match.match_date = row.day
        winner_id = match.player2_id if found.flipped else match.player1_id
        if match.winner_id is not None and match.winner_id != winner_id:
            self.report.add_warning(
                "ingest_tennis_data.warnings.winner_mismatch",
                path=self.path,
                line=row.line,
                winner=row.winner,
                match_id=match.id,
            )
        for (bookmaker, side), odds in row.odds.items():
            selection = 3 - side if found.flipped else side
            self.odds[(match.id, selection, bookmaker)] = (odds, row.day)

    # --- свои матчи ---

    async def _own(
        self,
        atp: int,
        group: list[TdRow],
        info: _GroupInfo,
        near: list[_Candidate],
        *,
        ambiguous: bool,
    ) -> None:
        if ambiguous:
            self.report.add_warning(
                "ingest_tennis_data.warnings.tournament_ambiguous",
                path=self.path,
                tournament=info.name,
            )
        labels = [row.round_label for row in group]
        first_round = sum(1 for label in labels if rounds.round_number(label) == 1)
        draw = rounds.tennis_data_draw_size(labels, first_round)
        ready: list[tuple[TdRow, str, int | None, int | None]] = []
        unresolved: list[TdRow] = []
        for row in group:
            winner, loser = self.index.name(row.winner), self.index.name(row.loser)
            in_sackmann = [
                item.match.id
                for cand in near
                if cand.covers(row.day)
                for item in cand.find(winner, loser, self.index, self.claimed)
            ]
            if in_sackmann:
                kind: UnmatchedKind = "tournament_ambiguous" if ambiguous else "in_sackmann"
                self._unmatched(kind, row, match_ids=sorted(in_sackmann))
                unresolved.append(row)
                continue
            try:
                code = rounds.tennis_data_round(row.round_label, draw)
            except ValueError as exc:
                self.report.skip("bad_round")
                self.report.add_warning(
                    "ingest_tennis_data.warnings.bad_round",
                    path=self.path,
                    line=row.line,
                    error=str(exc),
                )
                unresolved.append(row)
                continue
            winner_id = await self.index.resolve(self.db, winner, row.day)
            loser_id = await self.index.resolve(self.db, loser, row.day)
            ambiguous_ids = [
                choice.ids for choice in (winner_id, loser_id) if isinstance(choice, _Ambiguous)
            ]
            if ambiguous_ids:
                self._unmatched("player_ambiguous", row, candidates=ambiguous_ids)
                unresolved.append(row)
                continue
            assert not isinstance(winner_id, _Ambiguous) and not isinstance(loser_id, _Ambiguous)
            ready.append((row, code, winner_id, loser_id))
        tournament = self.own_tournaments.get(f"{self.year}-{atp}")
        if ready:
            tournament = await self._own_tournament(atp, info)
        if tournament is None:
            return
        self.own_groups.setdefault(tournament.id, []).extend(unresolved)
        for row, code, maybe_winner, maybe_loser in ready:
            winner_id = await self._player(maybe_winner, row.winner, row)
            loser_id = await self._player(maybe_loser, row.loser, row)
            if winner_id == loser_id:
                self._unmatched("same_player", row, player_id=winner_id)
                self.own_groups[tournament.id].append(row)
                continue
            ref = own_match_ref(self.year, atp, row.round_label, row.winner, row.loser)
            self.plans.append(_Plan(row, tournament.id, code, winner_id, loser_id, ref))

    async def _own_tournament(self, atp: int, info: _GroupInfo) -> Tournament:
        ref = f"{self.year}-{atp}"
        if info.series and info.category is None:
            self.report.add_warning(
                "ingest_tennis_data.warnings.series_unknown",
                path=self.path,
                tournament=info.name,
                series=info.series,
            )
        values: dict[str, Any] = {
            "name": info.name,
            "city": info.city,
            "surface": info.surface,
            "indoor": info.indoor,
            "category": info.category,
            "best_of": info.best_of,
            "start_date": info.start,
        }
        tournament = self.own_tournaments.get(ref)
        if tournament is None:
            tournament = Tournament(source=SOURCE, source_ref=ref, **values)
            self.db.add(tournament)
            await self.db.flush()
            self.own_tournaments[ref] = tournament
            self._count("tournaments.created", created=True)
            return tournament
        changed = False
        for key, value in values.items():
            if getattr(tournament, key) != value:
                setattr(tournament, key, value)
                changed = True
        if changed:
            self._count("tournaments.updated", created=False)
        return tournament

    async def _player(self, found: int | None, raw: str, row: TdRow) -> int:
        """id игрока; None — создать игрока «И. Фамилия» с алиасом «Фамилия И.»."""
        if found is not None:
            return found
        # Тот же новый игрок мог быть создан строкой выше
        again = await self.index.resolve(self.db, self.index.name(raw), row.day)
        if isinstance(again, int):
            return again
        player = Player(name=player_display_name(raw)[:200], aliases=[raw])
        self.db.add(player)
        await self.db.flush()
        self.index.add(player.id, [player.name, raw], searchable=True)
        self.index.years[player.id] = {row.day.year}
        self._count("players.created", created=True)
        self.report.add_warning(
            "ingest_tennis_data.warnings.player_created",
            name=player.name,
            source=raw,
            path=self.path,
            line=row.line,
        )
        return player.id

    # --- запись ---

    async def _write_match_dates(self) -> None:
        if not self.match_dates:
            return
        now = datetime.now(UTC)
        changes = [
            {"id": mid, "match_date": day, "updated_at": now}
            for mid, day in sorted(self.match_dates.items())
        ]
        for chunk in _batches(changes):
            await self.db.execute(update(Match), list(chunk))
        for _ in changes:
            self._count("matches.dated", created=False)

    async def _write_own_matches(self) -> None:
        """Свои матчи: найти в базе (_assign), лишние матчи своих турниров убрать
        (_remove_orphans), остальные обновить или вставить."""
        if not self.plans and not self.own_groups:
            return
        tournament_ids = sorted({plan.tournament_id for plan in self.plans} | set(self.own_groups))
        rows = (
            await self.db.execute(
                select(
                    Match.id,
                    Match.source,
                    Match.source_ref,
                    Match.tournament_id,
                    Match.round,
                    Match.player1_id,
                    Match.player2_id,
                    Match.winner_id,
                    Match.match_date,
                    Match.status,
                    Match.score,
                    Match.best_of,
                ).where(
                    or_(
                        Match.tournament_id.in_(tournament_ids),
                        and_(
                            Match.source == SOURCE,
                            Match.source_ref.in_(sorted({plan.ref for plan in self.plans})),
                        ),
                    )
                )
            )
        ).all()
        own = [row for row in rows if row.source == SOURCE]
        found = self._assign(own)
        used = {row.id for row in found.values()}
        removed = await self._remove_orphans(
            [row for row in own if row.tournament_id in self.own_groups and row.id not in used]
        )
        rows = [row for row in rows if row.id not in removed]
        # Ключ UNIQUE(tournament_id, round, player1_id, player2_id) занят любым матчем турнира,
        # UNIQUE(source, source_ref) — любым своим матчем
        taken = {(row.tournament_id, row.round, row.player1_id, row.player2_id) for row in rows}
        refs = {row.source_ref for row in rows if row.source == SOURCE}
        inserts: list[tuple[_Plan, dict[str, Any]]] = []
        updates: list[dict[str, Any]] = []
        ids: list[tuple[_Plan, int]] = []
        now = datetime.now(UTC)
        for number, plan in enumerate(self.plans):
            current = found.get(number)
            if current is None:
                if plan.key in taken or plan.ref in refs:
                    self._unmatched("duplicate_row", plan.row)
                    continue
                taken.add(plan.key)
                refs.add(plan.ref)
                inserts.append((plan, {**plan.values(), "source": SOURCE, "source_ref": plan.ref}))
                continue
            values = {**plan.values(), "source_ref": plan.ref}
            old_key = (current.tournament_id, current.round, current.player1_id, current.player2_id)
            if (plan.key != old_key and plan.key in taken) or (
                plan.ref != current.source_ref and plan.ref in refs
            ):
                self._unmatched("duplicate_row", plan.row, match_id=current.id)
                continue
            taken.discard(old_key)
            taken.add(plan.key)
            refs.discard(current.source_ref)
            refs.add(plan.ref)
            if any(getattr(current, key) != value for key, value in values.items()):
                updates.append({"id": current.id, **values, "updated_at": now})
            ids.append((plan, current.id))
        for new in _batches(inserts):
            result = await self.db.execute(
                insert(Match).values([values for _, values in new]).returning(Match.id)
            )
            for (plan, _), mid in zip(new, result.scalars().all(), strict=True):
                ids.append((plan, mid))
                self._count("matches.created", created=True)
        for changed in _batches(updates):
            await self.db.execute(update(Match), list(changed))
        for _ in updates:
            self._count("matches.updated", created=False)
        for plan, mid in ids:
            self.stats.own += 1
            self.report.counts["rows.own"] += 1
            self.index.played(plan.winner_id, plan.row.day.year)
            self.index.played(plan.loser_id, plan.row.day.year)
            for (bookmaker, side), odds in plan.row.odds.items():
                self.odds[(mid, side, bookmaker)] = (odds, plan.row.day)

    def _assign(self, own: Sequence[_DbRow]) -> dict[int, _DbRow]:
        """Номер плана → его матч в базе. Сначала по source_ref, затем по ключу UNIQUE,
        затем по паре «победитель — проигравший» в турнире, если она там одна и среди
        планов, и среди матчей (поменялась метка раунда, source_ref прежнего вида)."""
        found: dict[int, _DbRow] = {}
        used: set[int] = set()

        def take(number: int, row: _DbRow | None) -> None:
            if row is not None and row.id not in used:
                found[number] = row
                used.add(row.id)

        by_ref = {row.source_ref: row for row in own}
        for number, plan in enumerate(self.plans):
            take(number, by_ref.get(plan.ref))
        by_key = {
            (row.tournament_id, row.round, row.player1_id, row.player2_id): row for row in own
        }
        for number, plan in enumerate(self.plans):
            if number not in found:
                take(number, by_key.get(plan.key))
        free: defaultdict[tuple[int | None, int, int], list[_DbRow]] = defaultdict(list)
        for row in own:
            if row.id not in used:
                free[(row.tournament_id, row.player1_id, row.player2_id)].append(row)
        waiting = Counter(
            (plan.tournament_id, plan.winner_id, plan.loser_id)
            for number, plan in enumerate(self.plans)
            if number not in found
        )
        for number, plan in enumerate(self.plans):
            pair = (plan.tournament_id, plan.winner_id, plan.loser_id)
            if number not in found and waiting[pair] == 1 and len(free[pair]) == 1:
                take(number, free[pair][0])
        return found

    async def _remove_orphans(self, orphans: Sequence[_DbRow]) -> set[int]:
        """Свои матчи турниров этого файла, которых в файле нет. Матч удаляется (кэфы
        архива — каскадом), если на него ничего не ссылается и ему не может соответствовать
        строка группы, которая в этот раз не сопоставилась; иначе — предупреждение.
        Возвращает id удаленных."""
        if not orphans:
            return set()
        players = {pid for row in orphans for pid in (row.player1_id, row.player2_id)}
        unknown = players - self.index.keys.keys()
        if unknown:
            await self.index.load(self.db, unknown)
        names: dict[int, str] = {
            pid: name
            for pid, name in await self.db.execute(
                select(Player.id, Player.name).where(Player.id.in_(sorted(players)))
            )
        }
        tournaments = {tour.id: tour.name for tour in self.own_tournaments.values()}
        referenced = await self._referenced(sorted(row.id for row in orphans))
        removed: list[int] = []
        for row in orphans:
            keep = row.id in referenced or self._may_be_unmatched_row(row)
            self.report.add_warning(
                "ingest_tennis_data.warnings.own_match_kept"
                if keep
                else "ingest_tennis_data.warnings.own_match_removed",
                path=self.path,
                match_id=row.id,
                tournament=tournaments.get(row.tournament_id, row.tournament_id),
                round=row.round,
                player1=names.get(row.player1_id, row.player1_id),
                player2=names.get(row.player2_id, row.player2_id),
            )
            if not keep:
                removed.append(row.id)
        for chunk in _batches(removed):
            await self.db.execute(delete(Match).where(Match.id.in_(list(chunk))))
        for _ in removed:
            self._count("matches.deleted", created=False)
        return set(removed)

    def _may_be_unmatched_row(self, match: _DbRow) -> bool:
        """Пара матча похожа на строку его группы, которая не стала матчем (игрок
        неоднозначен и т. п.): матч, скорее всего, верный, удалять его нельзя."""
        first = self.index.player_keys(match.player1_id)
        second = self.index.player_keys(match.player2_id)
        for row in self.own_groups.get(match.tournament_id, []):
            winner, loser = self.index.name(row.winner), self.index.name(row.loser)
            if min(winner.level(first), loser.level(second)) >= 1:
                return True
            if min(winner.level(second), loser.level(first)) >= 1:
                return True
        return False

    async def _referenced(self, ids: list[int]) -> set[int]:
        """Матчи, на которые ссылается что-то кроме кэфов архива: ставки, прогнозы,
        досье, статистика. Внешние ключи на matches.id берутся из метаданных: новая
        таблица учтется сама."""
        found: set[int] = set()
        target = Match.__table__.c.id
        for table in Base.metadata.sorted_tables:
            for fk in table.foreign_keys:
                if fk.column is not target:
                    continue
                stmt = select(fk.parent).where(fk.parent.in_(ids)).distinct()
                if table is OddsSnapshot.__table__:
                    stmt = stmt.where(OddsSnapshot.kind != ODDS_KIND)
                found.update(await self.db.scalars(stmt))
        return found

    async def _write_odds(self) -> None:
        values = [
            {
                "match_id": mid,
                "market": MARKET,
                "selection": selection,
                "bookmaker": bookmaker,
                "odds": odds,
                "kind": ODDS_KIND,
                "taken_at": datetime.combine(day, time(0), tzinfo=UTC),
            }
            for (mid, selection, bookmaker), (odds, day) in sorted(self.odds.items())
        ]
        for chunk in _batches(values):
            base = insert(OddsSnapshot).values(list(chunk))
            stmt = base.on_conflict_do_update(
                index_elements=["match_id", "market", "selection", "bookmaker"],
                index_where=text("kind = 'archive'"),
                set_={"odds": base.excluded.odds, "taken_at": base.excluded.taken_at},
                where=or_(
                    OddsSnapshot.odds != base.excluded.odds,
                    OddsSnapshot.taken_at != base.excluded.taken_at,
                ),
            ).returning(literal_column("xmax = 0", Boolean))
            for inserted in (await self.db.execute(stmt)).scalars():
                self._count("odds.created" if inserted else "odds.updated", created=bool(inserted))


# --- загрузка ---


def _this_year() -> int:
    return datetime.now(UTC).year


def file_path(year: int) -> str:
    return f"{year}/{year}.xlsx"


def _forbidden(exc: FetchError) -> bool:
    """Сайт ответил 403: Cloudflare не пускает программы."""
    return exc.status == 403


async def _last_date_in_db(db: AsyncSession, year: int) -> date | None:
    """Последний день матча года в базе: свежесть для файла, который не изменился."""
    value = await db.scalar(
        select(func.max(Match.match_date)).where(
            Match.match_date.between(date(year, 1, 1), date(year, 12, 31))
        )
    )
    return value if isinstance(value, date) else None


async def _load(
    *,
    years: Sequence[int],
    source: DataSource,
    force: bool,
    db_factory: Callable[[], AsyncSession],
) -> IngestReport:
    settings = get_settings()
    this_year = _this_year()
    async with ingestion_run(db_factory, SOURCE) as run:
        report = run.report
        async with db_factory() as db:
            known = await last_ok_file_hashes(db, SOURCE)
            index = PlayerIndex()
            await index.load(db)
        per_year: dict[str, dict[str, Any]] = {}
        missing: list[str] = []
        # Файлы прошлых лет, которых нет на сайте: они там есть всегда, значит адрес неверен
        absent: list[str] = []
        found = 0
        remote = source.directory is None
        for year in sorted(set(years)):
            if not settings.tennis_data_first_year <= year <= this_year:
                report.add_warning(
                    "ingest_tennis_data.warnings.year_unsupported",
                    year=year,
                    first=settings.tennis_data_first_year,
                )
                continue
            path = file_path(year)
            try:
                file = await source.fetch(path)
            except FetchError as exc:
                if _forbidden(exc):
                    raise IngestError(
                        "ingest_tennis_data.errors.blocked", path=path, page=MANUAL_PAGE
                    ) from None
                raise
            if file is None:
                missing.append(path)
                report.counts["files.missing"] += 1
                per_year[str(year)] = {"status": "missing"}
                if year != this_year:
                    report.add_warning("ingest_tennis_data.warnings.file_missing", path=path)
                    if remote:
                        absent.append(path)
                continue
            found += 1
            if not force and known.get(path) == file.sha256:
                report.add_file(file)
                report.counts["files.unchanged"] += 1
                async with db_factory() as db:
                    last = await _last_date_in_db(db, year)
                per_year[str(year)] = {"status": "unchanged", "last_date": last}
                continue
            rows = await asyncio.to_thread(
                parse_file, file.content, path=path, year=year, report=report
            )
            async with db_factory() as db:
                stats = await _YearLoader(
                    db, year=year, path=path, rows=rows, index=index, report=report
                ).run()
                await db.commit()
            report.add_file(file)
            report.counts["files.loaded"] += 1
            per_year[str(year)] = {
                "status": "loaded",
                "rows": stats.rows,
                "first_date": stats.first_date,
                "last_date": stats.last_date,
                "pinnacle_share": round(stats.pinnacle / stats.rows, 4) if stats.rows else None,
                "linked": stats.linked,
                "own": stats.own,
                "unmatched": stats.unmatched,
            }
        dates = [info["last_date"] for info in per_year.values() if info.get("last_date")]
        report.freshness = {
            "years": per_year,
            "missing_files": missing,
            "last_match_date": max(dates) if dates else None,
            "current_year_last_date": per_year.get(str(this_year), {}).get("last_date"),
        }
        # Иначе запуск был бы «успешным» с нулем строк, и уведомление о старых данных
        # (data_stale) не пришло бы никогда. Файла текущего года может еще не быть
        if absent:
            raise IngestError(
                "ingest_tennis_data.errors.files_missing",
                files=", ".join(absent),
                location=source.location,
            )
        if not found and any(path != file_path(this_year) for path in missing):
            raise IngestError("ingest_tennis_data.errors.nothing_found", location=source.location)
        return report


async def load_tennis_data(
    *,
    years: Sequence[int],
    source: DataSource,
    force: bool = False,
    db_factory: Callable[[], AsyncSession] | None = None,
) -> IngestReport:
    """Загрузка файлов tennis-data за годы (см. docstring модуля), журнал —
    ingestion_runs(source='tennis_data'). db_factory — фабрика сеансов БД; без нее
    создается свой движок по настройкам. Файл не получен или не разобран — IngestError
    (запуск помечается failed)."""
    if db_factory is not None:
        return await _load(years=years, source=source, force=force, db_factory=db_factory)
    engine = create_engine()
    try:
        return await _load(
            years=years, source=source, force=force, db_factory=create_sessionmaker(engine)
        )
    finally:
        await engine.dispose()


# --- команда терминала ---


def summary_lines(report: IngestReport, location: str) -> list[str]:
    """Отчет загрузки для терминала."""
    counts = report.counts
    unmatched = sum(n for key, n in counts.items() if key.startswith("unmatched."))
    lines = [
        t("ingest_tennis_data.summary.title", location=location),
        t(
            "ingest_tennis_data.summary.files",
            loaded=counts["files.loaded"],
            unchanged=counts["files.unchanged"],
            missing=counts["files.missing"],
        ),
        t(
            "ingest_tennis_data.summary.rows",
            read=report.rows_read,
            linked=counts["rows.linked"],
            own=counts["rows.own"],
            skipped=report.skipped,
            unmatched=unmatched,
        ),
        t(
            "ingest_tennis_data.summary.written",
            created=report.created,
            updated=report.updated,
            matches=counts["matches.created"],
            matches_updated=counts["matches.updated"] + counts["matches.dated"],
            matches_deleted=counts["matches.deleted"],
            tournaments=counts["tournaments.created"],
            tournaments_updated=counts["tournaments.updated"],
            players=counts["players.created"],
            odds=counts["odds.created"],
            odds_updated=counts["odds.updated"],
        ),
    ]
    for year, info in sorted(report.freshness.get("years", {}).items()):
        status = info.get("status")
        last = info.get("last_date") or "—"
        if status == "loaded":
            share = info.get("pinnacle_share")
            pinnacle = f"{share:.0%}" if isinstance(share, float | int) else "—"
            lines.append(
                t(
                    "ingest_tennis_data.summary.year",
                    year=year,
                    rows=info["rows"],
                    last=last,
                    pinnacle=pinnacle,
                )
            )
        elif status == "unchanged":
            lines.append(t("ingest_tennis_data.summary.year_unchanged", year=year, last=last))
        elif int(year) < _this_year():
            lines.append(t("ingest_tennis_data.summary.year_absent", year=year))
        else:
            lines.append(t("ingest_tennis_data.summary.year_missing", year=year))
    if report.warnings:
        total = sum(n for key, n in counts.items() if key.startswith("warnings."))
        lines.append(t("ingest_tennis_data.summary.warnings", count=total))
        lines += [f"  {warning}" for warning in report.warnings[:20]]
        if total > 20:
            lines.append(t("ingest_tennis_data.summary.more", count=total - 20))
    if report.unmatched:
        lines.append(t("ingest_tennis_data.summary.unmatched", count=unmatched))
        for example in report.unmatched[:20]:
            lines.append(
                t(
                    "ingest_tennis_data.summary.unmatched_row",
                    file=example.get("file"),
                    line=example.get("line"),
                    winner=example.get("winner"),
                    loser=example.get("loser"),
                    date=example.get("date"),
                    tournament=example.get("tournament"),
                    reason=t(f"ingest_tennis_data.unmatched.{example.get('kind')}"),
                )
            )
        if unmatched > 20:
            lines.append(t("ingest_tennis_data.summary.more", count=unmatched - 20))
    return lines


def parse_years(value: str, *, first: int, last: int) -> list[int]:
    """«2024», «2013-2026», «2019,2024-2026» → список лет. Вне [first, last] — ValueError."""
    years: set[int] = set()
    for part in value.split(","):
        start, sep, end = part.strip().partition("-")
        try:
            low = int(start)
            high = int(end) if sep else low
        except ValueError:
            raise ValueError(value) from None
        if low > high or low < first or high > last:
            raise ValueError(value)
        years.update(range(low, high + 1))
    return sorted(years)


def build_parser(first: int, last: int) -> argparse.ArgumentParser:
    def years(value: str) -> list[int]:
        try:
            return parse_years(value, first=first, last=last)
        except ValueError:
            message = t("ingest_tennis_data.errors.years", value=value, first=first, last=last)
            raise argparse.ArgumentTypeError(message) from None

    parser = argparse.ArgumentParser(
        prog="python -m app.ingest.tennis_data", description=t("ingest_tennis_data.description")
    )
    parser.add_argument(
        "--years", type=years, help=t("ingest_tennis_data.help.years", first=first, last=last)
    )
    parser.add_argument("--dir", help=t("ingest_tennis_data.help.dir"))
    parser.add_argument("--force", action="store_true", help=t("ingest_tennis_data.help.force"))
    return parser


async def _run(years: list[int], directory: str | None, force: bool) -> tuple[IngestReport, str]:
    source = (
        DataSource.local(directory)
        if directory is not None
        else DataSource.http(get_settings().tennis_data_base_url)
    )
    engine = create_engine()
    try:
        factory = create_sessionmaker(engine)
        # Не одновременно с ночным обновлением и другими командами шагов (runs.update_lock)
        async with update_lock(factory), source:
            report = await load_tennis_data(
                years=years, source=source, force=force, db_factory=factory
            )
            return report, source.location
    finally:
        await engine.dispose()


def main(argv: Sequence[str] | None = None) -> int:
    """Код возврата: 0 — загрузка прошла (возможно, с замечаниями), 1 — ошибка."""
    first, last = get_settings().tennis_data_first_year, _this_year()
    args = build_parser(first, last).parse_args(argv)
    if args.dir is not None and not Path(args.dir).is_dir():
        print(t("ingest_tennis_data.errors.dir_not_found", path=args.dir), file=sys.stderr)
        return 1
    years: list[int] = args.years or list(range(first, last + 1))
    try:
        report, location = asyncio.run(_run(years, args.dir, args.force))
    except IngestError as exc:
        print(exc.message, file=sys.stderr)
        return 1
    except (DBAPIError, OSError):
        # DBAPIError — общий класс ошибок драйвера: нет связи, неверный пароль, нет базы
        # или таблиц. Текст исключения не печатаем: в нем бывает адрес базы с паролем
        print(t("ingest_tennis_data.errors.db_unavailable"), file=sys.stderr)
        return 1
    print("\n".join(summary_lines(report, location)), flush=True)
    return 0


if __name__ == "__main__":
    sys.exit(main())
